"""Export a PPTX whose manim slides embed the *playable* MP4 video.

Slidev's own `export --format pptx` rasterises every slide to an image, so a
manim slide becomes a still frame — the animation is lost. This builder instead:

  1. renders one PNG per slide (final click state) via `slidev export --format png`,
  2. converts each manim scene's webm → H.264 MP4 (the only codec PowerPoint
     embeds reliably),
  3. assembles a .pptx where ordinary slides are the PNG and manim slides embed
     the MP4 full-bleed (poster = the last-frame image from build-manim),

so the deck opens in PowerPoint/Keynote with the manim scenes playing in place.

Requirements: Node + Slidev (for the PNG export), ffmpeg (webm→mp4), and
`python-pptx` (`pip install presentation-sanity[pptx]`).
"""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

import yaml

from .manifest import Manifest
from .manim_render import ffmpeg_path, poster_path_for

# 16:9 slide, in EMU (English Metric Units; 914400 per inch).
_EMU_PER_INCH = 914400
_SLIDE_W = int(13.333 * _EMU_PER_INCH)
_SLIDE_H = int(7.5 * _EMU_PER_INCH)


# ── slide parsing (mirrors @slidev/parser parseSync + the hide/disabled filter)

def _split_slides(markdown: str) -> list[str]:
    """Split deck markdown into per-slide raw blocks, matching Slidev's parser.

    Separators are lines beginning with `---`; a `---` immediately followed by a
    non-blank line opens a frontmatter block that runs to the next `---` (so the
    fences inside frontmatter aren't mistaken for separators). Fenced code blocks
    are skipped so a `---` inside ``` doesn't split a slide.
    """
    lines = markdown.replace("\r\n", "\n").split("\n")
    n = len(lines)
    slides: list[str] = []
    start = 0

    def emit(end: int) -> None:
        nonlocal start
        if start != end:
            slides.append("\n".join(lines[start:end]))
        start = end + 1

    i = 0
    while i < n:
        line = lines[i].rstrip()
        if line.startswith("---"):
            emit(i)
            nxt = lines[i + 1] if i + 1 < n else None
            # `---` (not `----`) followed by content → frontmatter block
            if (len(line) <= 3 or line[3] != "-") and (nxt is not None and nxt.strip()):
                start = i
                i += 1
                while i < n and lines[i].rstrip() != "---":
                    i += 1
        elif line.lstrip().startswith("```"):
            fence = re.match(r"^\s*`+", line).group(0)
            j = i + 1
            while j < n and not lines[j].startswith(fence.lstrip()):
                j += 1
            if j != n:
                i = j
        i += 1
    if start <= n - 1:
        emit(n)
    return slides


def _frontmatter(raw: str) -> dict[str, Any]:
    """Parse a slide block's leading YAML frontmatter (`---\\n…\\n---`)."""
    if not raw.lstrip().startswith("---"):
        return {}
    m = re.match(r"^\s*---\n(.*?)\n---", raw, re.S)
    if not m:
        return {}
    try:
        data = yaml.safe_load(m.group(1))
        return data if isinstance(data, dict) else {}
    except yaml.YAMLError:
        return {}


def parse_slide_frontmatters(markdown: str) -> list[dict[str, Any]]:
    """Frontmatter dicts for every *rendered* slide, in order.

    Slides with `hide:`/`disabled:` truthy are dropped — Slidev excludes them, so
    this keeps our indices aligned with the PNG export's 1..N numbering.
    """
    out: list[dict[str, Any]] = []
    for raw in _split_slides(markdown):
        fm = _frontmatter(raw)
        if fm.get("hide") or fm.get("disabled"):
            continue
        out.append(fm)
    return out


# ── ffmpeg: webm → mp4 (H.264) ──────────────────────────────────────────────

def _ensure_mp4(webm: Path, mp4: Path, *, verbose: bool = False) -> bool:
    """Transcode `webm` → `mp4` (H.264/yuv420p, faststart, no audio). Cached by mtime."""
    if mp4.is_file() and mp4.stat().st_mtime >= webm.stat().st_mtime:
        return True
    ff = ffmpeg_path()
    if ff is None:
        return False
    mp4.parent.mkdir(parents=True, exist_ok=True)
    cmd = [
        ff, "-y",
        "-i", str(webm),
        "-c:v", "libx264",
        "-pix_fmt", "yuv420p",
        "-movflags", "+faststart",
        "-an",
        str(mp4),
    ]
    if verbose:
        print(f"  $ {' '.join(cmd)}", file=sys.stderr)
    proc = subprocess.run(
        cmd,
        stdout=subprocess.DEVNULL,
        stderr=None if verbose else subprocess.DEVNULL,
        check=False,
    )
    return proc.returncode == 0 and mp4.is_file()


# ── slidev PNG export (one image per slide, final click state) ───────────────

def _export_pngs(
    root: Path, out_dir: Path, *, with_clicks: bool = True, verbose: bool = False
) -> list[dict[str, Any]]:
    """Run `slidev export --format png` into `out_dir`; return ordered steps.

    Each step is `{"slide": int, "step": int, "path": Path}`. With `--with-clicks`
    Slidev emits one PNG per click state, named `<slide:03d>-<step:02d>.png`, so a
    build slide becomes several steps (its v-clicks revealed progressively). Manim
    slides have no clicks → a single `<slide>-01.png`. Without clicks, files are
    `<n>.png` (final state only); we normalise those to step 1.
    """
    if shutil.which("npx") is None:
        raise RuntimeError("npx not found on PATH. Install Node.js >= 20.")
    out_dir.mkdir(parents=True, exist_ok=True)
    cmd = ["npx", "slidev", "export", "--format", "png", "--output", str(out_dir)]
    if with_clicks:
        cmd.append("--with-clicks")
    if verbose:
        print(f"  $ (cwd={root}) {' '.join(cmd)}", file=sys.stderr)
    subprocess.run(
        cmd, cwd=root, check=True,
        stdout=None if verbose else subprocess.DEVNULL,
    )
    steps: list[dict[str, Any]] = []
    for p in out_dir.glob("*.png"):
        m = re.fullmatch(r"(\d+)-(\d+)", p.stem)
        if m:
            steps.append({"slide": int(m.group(1)), "step": int(m.group(2)), "path": p})
        elif p.stem.isdigit():
            steps.append({"slide": int(p.stem), "step": 1, "path": p})
    if not steps:
        raise RuntimeError(f"slidev produced no PNGs in {out_dir}")
    steps.sort(key=lambda s: (s["slide"], s["step"]))
    return steps


# ── pptx assembly ────────────────────────────────────────────────────────────

def _set_autoplay(slide) -> None:  # type: ignore[no-untyped-def]
    """Make the embedded movie start automatically on slide entry.

    `add_movie` builds a `<p:video>` media node whose start condition is
    `<p:cond delay="indefinite"/>` — i.e. wait for a click. Setting that delay to
    `0` (and dropping any `evt`) starts the video as soon as the slide is shown.
    Best effort: harmless if the timing structure differs.
    """
    from pptx.oxml.ns import qn

    timing = slide._element.find(qn("p:timing"))
    if timing is None:
        return
    for video in timing.iter(qn("p:video")):
        for cond in video.iter(qn("p:cond")):
            cond.attrib.pop("evt", None)
            cond.set("delay", "0")


def build_pptx(
    *,
    steps: list[dict[str, Any]],
    manim_slides: dict[int, str],
    mp4_for: dict[str, Path],
    poster_for: dict[str, Path],
    out_path: Path,
    autoplay: bool = True,
) -> int:
    """Assemble the PPTX. One slide per click step; manim slides embed the video
    once (extra steps of a manim slide, if any, are skipped). Returns slide count."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width = Emu(_SLIDE_W)
    prs.slide_height = Emu(_SLIDE_H)
    blank = prs.slide_layouts[6]  # fully blank layout

    embedded: set[int] = set()
    count = 0
    for st in steps:
        no = st["slide"]
        scene = manim_slides.get(no)
        mp4 = mp4_for.get(scene) if scene else None

        # A manim slide is a single video slide — skip any further click steps.
        if scene and mp4 and no in embedded:
            continue

        slide = prs.slides.add_slide(blank)
        count += 1

        if scene and mp4 and mp4.is_file():
            # Black background so letterboxing matches the scene's backdrop.
            bg = slide.background
            bg.fill.solid()
            bg.fill.fore_color.rgb = RGBColor(0, 0, 0)
            poster = poster_for.get(scene)
            slide.shapes.add_movie(
                str(mp4),
                Emu(0), Emu(0), Emu(_SLIDE_W), Emu(_SLIDE_H),
                poster_frame_image=str(poster) if poster and poster.is_file() else None,
                mime_type="video/mp4",
            )
            if autoplay:
                _set_autoplay(slide)
            embedded.add(no)
        else:
            slide.shapes.add_picture(
                str(st["path"]), Emu(0), Emu(0), Emu(_SLIDE_W), Emu(_SLIDE_H)
            )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return count


# ── orchestration ────────────────────────────────────────────────────────────

def export_pptx(
    root: Path,
    *,
    out: str | None = None,
    autoplay: bool = True,
    with_clicks: bool = True,
    verbose: bool = False,
) -> Path:
    """Build a video-embedding PPTX for the deck at `root`. Returns the output path."""
    try:
        import pptx  # noqa: F401
    except ModuleNotFoundError as e:
        raise RuntimeError(
            "python-pptx is required for export-pptx.\n"
            "  install with: pip install presentation-sanity[pptx]"
        ) from e

    manifest_md = root / "slides.md"
    if not manifest_md.is_file():
        raise RuntimeError(f"slides.md not found at {manifest_md}")

    manifest = _load_manifest_safe(root)
    frontmatters = parse_slide_frontmatters(manifest_md.read_text())
    n_slides = len(frontmatters)

    # which slides are manim, and to which scene
    manim_slides: dict[int, str] = {}
    for idx, fm in enumerate(frontmatters, start=1):
        if fm.get("layout") == "manim" and fm.get("scene"):
            manim_slides[idx] = str(fm["scene"])
    print(f"  parsed {n_slides} slides, {len(manim_slides)} manim")

    # transcode the manim scenes we need → mp4 (+ collect posters)
    if manim_slides and ffmpeg_path() is None:
        print(
            "  warning: ffmpeg not found — manim slides will fall back to static "
            "images (no embedded video)."
        )
    mp4_dir = root / ".cache" / "pptx-mp4"
    mp4_for: dict[str, Path] = {}
    poster_for: dict[str, Path] = {}
    for scene_key in sorted(set(manim_slides.values())):
        scene = manifest.scenes.get(scene_key) if manifest else None
        fmt = scene.format if scene else "webm"
        webm = root / "public" / "manim" / f"{scene_key}.{fmt}"
        if not webm.is_file():
            print(f"  warning: video for scene {scene_key!r} not found at {webm}")
            continue
        mp4 = mp4_dir / f"{scene_key}.mp4"
        if _ensure_mp4(webm, mp4, verbose=verbose):
            mp4_for[scene_key] = mp4
            print(f"  {scene_key}: mp4")
        poster = poster_path_for(webm)
        if poster.is_file():
            poster_for[scene_key] = poster

    # render PNGs — one per click step (so v-clicks become progressive slides)
    with tempfile.TemporaryDirectory(prefix="psanity-pptx-") as tmp:
        png_dir = Path(tmp) / "png"
        mode = "with click steps" if with_clicks else "final state only"
        print(f"  rendering slide PNGs ({mode})...")
        steps = _export_pngs(root, png_dir, with_clicks=with_clicks, verbose=verbose)
        rendered_slides = {s["slide"] for s in steps}
        if rendered_slides and max(rendered_slides) != n_slides:
            print(
                f"  note: parsed {n_slides} slides but slidev rendered "
                f"{max(rendered_slides)} — using the rendered numbering."
            )

        deck_name = (manifest.metadata.get("title") if manifest else None) or root.name
        safe = re.sub(r"[^A-Za-z0-9._-]+", "-", str(deck_name)).strip("-") or "deck"
        out_path = Path(out) if out else (root / "exports" / f"{safe}.pptx")

        print(f"  assembling {out_path.name} ({len(steps)} steps) ...")
        n_out = build_pptx(
            steps=steps,
            manim_slides=manim_slides,
            mp4_for=mp4_for,
            poster_for=poster_for,
            out_path=out_path,
            autoplay=autoplay,
        )

    print(f"  done → {out_path}")
    print(
        f"  {n_out} slides — {len(mp4_for)} embedded manim video(s), "
        f"{n_out - len(mp4_for)} image slide(s)."
    )
    return out_path


def _load_manifest_safe(root: Path) -> Manifest | None:
    """Load the manifest if present; manim slides still work via convention if not."""
    from .manifest import ManifestError, load_manifest

    try:
        return load_manifest(root)
    except ManifestError:
        return None
